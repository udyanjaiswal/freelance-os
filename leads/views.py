
from django.shortcuts import get_object_or_404, redirect, render
from django.http import HttpResponse
from django.contrib import messages
from django.core.validators import URLValidator
from django.core.exceptions import ValidationError

from .models import Lead
from outreach.models import Outreach
from scoring.engine import get_score_breakdown
from discovery.engine import save_lead_data

import csv
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import io


from openpyxl import load_workbook
from django.contrib.auth.decorators import login_required

@login_required
def lead_detail(request, lead_id):
    lead = get_object_or_404(
        Lead,
        id=lead_id,
        user=request.user,
    )

    # Save notes
    if request.method == "POST":
        notes = request.POST.get("notes")

        if notes is not None:
            lead.notes = notes
            lead.save(
                update_fields=[
                    "notes",
                    "updated_at",
                ]
            )

    score, breakdown = get_score_breakdown(lead)

    outreach_records = lead.outreach_records.all().order_by(
        "-created_at"
    )

    return render(
        request,
        "leads/lead_detail.html",
        {
            "lead": lead,
            "score": score,
            "breakdown": breakdown,
            "outreach_records": outreach_records,
        },
    )

@login_required
def update_status(request, lead_id):
    lead = get_object_or_404(
        Lead,
        id=lead_id,
        user=request.user,
    )

    if request.method == "POST":
        status = request.POST.get("status")

        valid_statuses = [
            choice[0]
            for choice in Lead.STATUS_CHOICES
        ]

        if status in valid_statuses:
            lead.status = status
            lead.save(
                update_fields=[
                    "status",
                    "updated_at",
                ]
            )
        else:
            # B14: surface invalid status to user
            messages.error(
                request,
                f"'{status}' is not a valid status. "
                "Please choose from the available options."
            )

    return redirect("lead_detail", lead_id=lead.id)

@login_required
def add_outreach(request, lead_id):
    lead = get_object_or_404(Lead, id=lead_id,
        user=request.user,)

    if request.method == "POST":
        method = request.POST.get("method", "").strip()
        follow_up_raw = request.POST.get("follow_up_date", "").strip()

        # B6: validate outreach method
        valid_methods = [
            choice[0] for choice in Outreach.METHOD_CHOICES
        ]
        if method not in valid_methods:
            messages.error(
                request,
                f"'{method}' is not a valid outreach method. "
                f"Choose from: {', '.join(valid_methods)}."
            )
            return redirect("lead_detail", lead_id=lead.id)

        # V4: follow-up date must be today or in the future
        follow_up_date = follow_up_raw or None
        if follow_up_raw:
            try:
                parsed_date = date.fromisoformat(follow_up_raw)
                if parsed_date < date.today():
                    messages.error(
                        request,
                        "Follow-up date cannot be in the past."
                    )
                    return redirect("lead_detail", lead_id=lead.id)
                follow_up_date = follow_up_raw
            except ValueError:
                messages.error(
                    request,
                    "Invalid follow-up date format."
                )
                return redirect("lead_detail", lead_id=lead.id)

        Outreach.objects.create(
            lead=lead,
            method=method,
            message=request.POST.get("message", ""),
            outcome=request.POST.get("outcome", ""),
            follow_up_date=follow_up_date,
        )

        # First outreach automatically moves
        # a new lead to contacted.
        if lead.status == "new":
            lead.status = "contacted"
            lead.save(
                update_fields=[
                    "status",
                    "updated_at",
                ]
            )

    return redirect("lead_detail", lead_id=lead.id)


@login_required
def import_leads(request):
    if request.method != "POST":
        return redirect("home")

    uploaded_file = request.FILES.get("csv_file")

    if not uploaded_file:
        messages.error(request, "Please select a CSV or Excel file.")
        return redirect("home")

    file_ext = os.path.splitext(uploaded_file.name)[1].lower()

    if file_ext not in [".csv", ".xlsx"]:
        messages.error(request, "Only CSV and XLSX files are supported.")
        return redirect("home")

    # Actual Lead model fields
    allowed_fields = {
        "business_name",
        "category",
        "address",
        "city",
        "phone",
        "email",
        "website",
        "source",
        "source_id",
        "notes",
    }

    # Different possible column names -> actual model field
    aliases = {
        "business_name": [
            "business_name", "business name", "business",
            "company", "company_name", "company name",
            "name", "lead_name", "lead name",
        ],
        "category": ["category", "industry", "business_type", "type"],
        "address": ["address", "full_address", "full address", "street"],
        "city": ["city", "town", "location_city"],
        "phone": ["phone", "phone_number", "phone number", "mobile", "telephone"],
        "email": ["email", "email_address", "email address"],
        "website": ["website", "web", "url", "site"],
        "source": ["source", "lead_source", "lead source"],
        "source_id": ["source_id", "source id", "place_id", "place id"],
        "notes": ["notes", "note", "comments", "description"],
    }

    # Normalize header names
    def normalize_header(value):
        return str(value or "").strip().lower().replace("-", "_")

    alias_lookup = {}

    for field, names in aliases.items():
        for name in names:
            alias_lookup[normalize_header(name)] = field

    # D3: URL validator instance
    _url_validator = URLValidator()

    def is_valid_url(value):
        try:
            _url_validator(value)
            return True
        except ValidationError:
            return False

    new_count = 0
    duplicate_count = 0
    error_count = 0
    imported_count = 0

    workbook = None

    try:
        # Read CSV
        if file_ext == ".csv":
            try:
                content = uploaded_file.read().decode("utf-8-sig")
            except UnicodeDecodeError:
                uploaded_file.seek(0)
                content = uploaded_file.read().decode("latin-1")

            reader = csv.DictReader(io.StringIO(content))

            if not reader.fieldnames:
                messages.error(request, "CSV file is empty or has no headers.")
                return redirect("home")

            headers = reader.fieldnames
            rows = reader

        # Read Excel
        else:
            workbook = load_workbook(
                uploaded_file,
                read_only=True,
                data_only=True
            )

            sheet = workbook.active
            sheet_rows = sheet.iter_rows(values_only=True)

            headers = next(sheet_rows, None)

            if not headers:
                messages.error(request, "Excel file is empty.")
                return redirect("home")

            rows = (
                dict(zip(headers, row))
                for row in sheet_rows
            )

        # Map uploaded headers to actual model fields
        column_map = {}

        for header in headers:
            normalized = normalize_header(header)
            actual_field = alias_lookup.get(normalized)

            if actual_field and actual_field not in column_map.values():
                column_map[header] = actual_field

        # Business name is mandatory
        if "business_name" not in column_map.values():
            messages.error(
                request,
                "File must contain a business name column "
                "(e.g. business_name, company, or name)."
            )
            return redirect("home")

        # Import rows
        for row in rows:
            try:
                data = {}

                for uploaded_column, actual_field in column_map.items():
                    value = row.get(uploaded_column)

                    if value is not None:
                        value = str(value).strip()

                    if value:
                        # D3: skip invalid website URLs silently
                        if actual_field == "website" and not is_valid_url(value):
                            continue
                        data[actual_field] = value

                # Ignore completely empty rows
                if not data:
                    continue

                # Skip rows without business name
                if not data.get("business_name"):
                    error_count += 1
                    continue

                result = save_lead_data(data, user=request.user)

                # Adapt result counting to your existing helper's return format
                if result == "duplicate":
                    duplicate_count += 1
                else:
                    new_count += 1

                imported_count += 1

            except Exception:
                error_count += 1

        messages.success(
            request,
            f"Import complete: {new_count} new, "
            f"{duplicate_count} duplicates, "
            f"{error_count} errors."
        )

    except Exception as e:
        messages.error(request, f"Import failed: {str(e)}")

    finally:
        if workbook:
            workbook.close()

    return redirect("home")

@login_required
def export_leads(request):
    # B5: only export the current user's leads — never expose another user's data
    leads = Lead.objects.filter(
        user=request.user
    ).order_by("-created_at")

    fields = [
        ("Business Name", "business_name"),
        ("Category", "category"),
        ("Address", "address"),
        ("City", "city"),
        ("Phone", "phone"),
        ("Email", "email"),
        ("Website", "website"),
        ("Score", "score"),
        ("Potential", "potential"),
        ("Status", "status"),
        ("Notes", "notes"),
    ]

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Leads"

    # Report title
    worksheet.merge_cells("A1:K1")
    worksheet["A1"] = "FREELANCE OS — LEADS REPORT"
    worksheet["A1"].font = Font(
        bold=True,
        size=16,
        color="FFFFFF"
    )
    worksheet["A1"].fill = PatternFill(
        "solid",
        fgColor="163DA1"
    )
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )
    worksheet.row_dimensions[1].height = 32

    # Column headers
    for column, (label, field) in enumerate(fields, start=1):
        cell = worksheet.cell(row=3, column=column)
        cell.value = label
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="0F2F80"
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    worksheet.row_dimensions[3].height = 25

    # Lead records
    for row_number, lead in enumerate(leads, start=4):
        for column, (label, field) in enumerate(fields, start=1):
            cell = worksheet.cell(
                row=row_number,
                column=column
            )
            cell.value = getattr(lead, field, "") or ""
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            if row_number % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="F1F5F9"
                )

    # Column widths
    widths = [
        25, 18, 35, 18, 18, 28, 30,
        18, 28, 12, 15, 15, 35
    ]

    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # Freeze headers and enable filters
    worksheet.freeze_panes = "A4"
    worksheet.sheet_view.showGridLines = False

    if worksheet.max_row >= 4:
        worksheet.auto_filter.ref = (
            f"A3:M{worksheet.max_row}"
        )

    # Excel download response
    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="Freelance_OS_Leads.xlsx"'
    )

    workbook.save(response)

    return response