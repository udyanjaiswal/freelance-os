import io
import openpyxl
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
from leads.models import Lead


class LeadsSecurityTests(TestCase):
    def setUp(self):
        self.user_a = User.objects.create_user(
            username="user_a",
            email="user_a@example.com",
            password="Password123!",
        )
        self.user_b = User.objects.create_user(
            username="user_b",
            email="user_b@example.com",
            password="Password123!",
        )

        self.lead_a = Lead.objects.create(
            user=self.user_a,
            business_name="Alpha Tech",
            email="alpha@example.com",
            phone="1234567890",
            status="new",
        )
        self.lead_b = Lead.objects.create(
            user=self.user_b,
            business_name="Beta Corp",
            email="beta@example.com",
            phone="0987654321",
            status="new",
        )

        self.client_a = Client()
        self.client_a.login(username="user_a", password="Password123!")

        self.anon_client = Client()

    def test_unauthenticated_requests_redirect(self):
        """Unauthenticated requests to leads endpoints must redirect to login."""
        detail_url = reverse("lead_detail", args=[self.lead_a.id])
        response = self.anon_client.get(detail_url)
        self.assertRedirects(response, f"/accounts/login/?next={detail_url}")

    def test_idor_cross_user_detail_blocked(self):
        """User A attempting to view User B's lead must receive 404 Not Found."""
        url = reverse("lead_detail", args=[self.lead_b.id])
        response = self.client_a.get(url)
        self.assertEqual(response.status_code, 404)

    def test_idor_cross_user_status_update_blocked(self):
        """User A attempting to update User B's lead status must receive 404 Not Found."""
        url = reverse("update_status", args=[self.lead_b.id])
        response = self.client_a.post(url, {"status": "contacted"})
        self.assertEqual(response.status_code, 404)

        # Confirm User B's lead status was not modified
        self.lead_b.refresh_from_db()
        self.assertEqual(self.lead_b.status, "new")

    def test_idor_cross_user_add_outreach_blocked(self):
        """User A attempting to add outreach to User B's lead must receive 404 Not Found."""
        url = reverse("add_outreach", args=[self.lead_b.id])
        response = self.client_a.post(url, {
            "method": "call",
            "message": "Unauthorized outreach attempt",
        })
        self.assertEqual(response.status_code, 404)
        self.assertEqual(self.lead_b.outreach_records.count(), 0)

    def test_state_modifying_endpoints_reject_get(self):
        """State-modifying endpoints must reject GET requests with 405 Method Not Allowed."""
        status_url = reverse("update_status", args=[self.lead_a.id])
        self.assertEqual(self.client_a.get(status_url).status_code, 405)

        outreach_url = reverse("add_outreach", args=[self.lead_a.id])
        self.assertEqual(self.client_a.get(outreach_url).status_code, 405)

        import_url = reverse("import_leads")
        self.assertEqual(self.client_a.get(import_url).status_code, 405)

    def test_export_leads_isolation(self):
        """Export must contain only the requesting user's leads."""
        export_url = reverse("export_leads")
        response = self.client_a.get(export_url)
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = wb.active
        exported_business_names = [
            sheet.cell(row=r, column=1).value
            for r in range(4, sheet.max_row + 1)
        ]

        self.assertIn("Alpha Tech", exported_business_names)
        self.assertNotIn("Beta Corp", exported_business_names)

    def test_formula_injection_sanitization_on_export(self):
        """Formulas in lead fields must be prepended with a quote to neutralize execution."""
        Lead.objects.create(
            user=self.user_a,
            business_name="=cmd|' /C calc'!A0",
            notes="+2+5",
            address="@SUM(1,2)",
        )

        export_url = reverse("export_leads")
        response = self.client_a.get(export_url)
        self.assertEqual(response.status_code, 200)

        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        sheet = wb.active

        # Find row for the malicious lead
        found_sanitized = False
        for r in range(4, sheet.max_row + 1):
            name_val = sheet.cell(row=r, column=1).value
            if name_val and "=cmd" in name_val:
                self.assertTrue(name_val.startswith("'="), f"Expected formula to be neutralized, got: {name_val}")
                found_sanitized = True

        self.assertTrue(found_sanitized, "Sanitized formula lead was not found in export")
